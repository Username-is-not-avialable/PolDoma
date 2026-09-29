"""Импорт контактов из Excel-файла в таблицу contacts.

Ожидаемый формат файла: столбец A — название компании / ФИО, столбец B — ИНН,
первая строка — заголовки таблицы.

Валидным ИНН считается только непустая последовательность цифр без каких-либо
других символов (пробелов, запятых, переносов строк, букв). Всё остальное —
пустые ячейки, несколько ИНН в одной ячейке, «------------», текстовые пометки
вида «удалить» — пропускается и попадает в итоговый отчёт.

Использование:
    python scripts/import_contacts.py                      # файл data/список исцов.xlsx
    python scripts/import_contacts.py --dry-run            # только отчёт, без записи в БД
    python scripts/import_contacts.py --file other.xlsx --sheet Лист1
    python scripts/import_contacts.py --clean              # очистить contacts перед импортом
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import re
import sys
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import async_session_factory, engine
from app.models.contact import Contact

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("import_contacts")

# Каталог проекта и путь по умолчанию к файлу с данными (data/ — в .gitignore,
# поэтому сам файл в репозиторий не попадает; путь не зависит от текущего каталога)
PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_FILE_REL = Path("data") / "список исцов.xlsx"
DEFAULT_FILE = PROJECT_ROOT / DEFAULT_FILE_REL
DEFAULT_EMAIL = "antonbakir280@gmail.com"

# Индексы столбцов листа (0-based): A — название, B — ИНН
COL_NAME = 0
COL_INN = 1

# Маркеры шапки таблицы (строка 1 не импортируется)
HEADER_NAME_MARKERS = ("истец", "наимен", "назван", "компан", "name", "company")
HEADER_INN_MARKERS = ("инн", "inn")

# Названия-пометки: такие строки импортировать не надо. Сравнение точное,
# чтобы не отбросить реальные компании вроде «Удалов И.И.»
SKIP_NAME_VALUES = frozenset({"удалить", "удали", "удалено", "delete"})

# Сколько записей отправлять одним INSERT (upsert для PostgreSQL)
CHUNK_SIZE = 500

# Валидный ИНН: только цифры, длина не ограничена
INN_RE = re.compile(r"\d+")

# Причины, по которым строка не попадает в БД
REASON_EMPTY = "пусто"
REASON_INVALID = "невалидно"


@dataclass
class ParseReport:
    """Результат разбора листа Excel (БД не затрагивается)."""

    sheet_name: str = ""
    total_rows: int = 0
    skipped_empty_inn: int = 0
    skipped_invalid_inn: int = 0
    skipped_skip_name: int = 0
    skipped_empty_name: int = 0
    skipped_duplicates: int = 0
    contacts: list[dict[str, Any]] = field(default_factory=list)
    invalid_samples: list[str] = field(default_factory=list)

    @property
    def imported(self) -> int:
        """Количество контактов, готовых к записи в БД."""
        return len(self.contacts)


def normalize_inn(value: Any) -> tuple[str | None, str | None]:
    """Приводит значение ячейки ИНН к строке цифр.

    Возвращает пару (инн, причина_пропуска): для валидного значения причина
    равна None. Валидны только непустые последовательности цифр любой длины
    без иных знаков; ведущие нули сохраняются (к числу не приводим).
    """
    if value is None:
        return None, REASON_EMPTY
    if isinstance(value, bool):  # bool — подкласс int, ИНН им быть не может
        return None, REASON_INVALID
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float):
        # Excel может отдать 10–12-значный ИНН как число с плавающей точкой
        if not value.is_integer():
            return None, REASON_INVALID
        text = str(int(value))
    else:
        text = str(value).replace("\xa0", " ").strip()
        if not text:
            return None, REASON_EMPTY
    if not INN_RE.fullmatch(text):
        return None, REASON_INVALID
    return text, None


def normalize_name(value: Any) -> str | None:
    """Приводит название компании/ФИО к строке без лишних пробелов и переносов."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    return text or None


def is_skippable_name(name: str) -> bool:
    """True, если название — служебная пометка (например, «удалить»)."""
    return name.strip().lower().rstrip(".!") in SKIP_NAME_VALUES


def validate_email(email: str) -> str:
    """Простая проверка адреса email (есть локальная часть, @ и домен)."""
    value = email.strip()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", value):
        raise ValueError(f"Некорректный email: {email!r}")
    return value


def _looks_like_header(name_raw: Any, inn_raw: Any) -> bool:
    """Эвристика: первая строка — шапка таблицы (например, «Истец» / «ИНН»)."""
    name = str(name_raw or "").strip().lower()
    inn = str(inn_raw or "").strip().lower()
    return any(m in name for m in HEADER_NAME_MARKERS) and any(
        m in inn for m in HEADER_INN_MARKERS
    )


def _shorten(value: Any, limit: int = 60) -> str:
    """Короткое представление значения для лога."""
    text = repr(value)
    return text if len(text) <= limit else text[: limit - 1] + "…"


def parse_workbook(
    path: Path,
    sheet: str | None = None,
    email: str = DEFAULT_EMAIL,
) -> ParseReport:
    """Читает Excel-файл и возвращает отчёт со списком валидных контактов.

    Только чтение: ни файл, ни БД не изменяются.
    """
    if not path.is_file():
        raise FileNotFoundError(f"Файл не найден: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None:
            if sheet not in workbook.sheetnames:
                raise ValueError(
                    f"Лист {sheet!r} не найден. Доступные листы: {workbook.sheetnames}"
                )
            worksheet = workbook[sheet]
        else:
            worksheet = workbook[workbook.sheetnames[0]]

        report = ParseReport(sheet_name=worksheet.title)
        _collect_rows(worksheet, report, email=email, source=path.name)
    finally:
        workbook.close()
    return report


def _collect_rows(worksheet: Any, report: ParseReport, email: str, source: str) -> None:
    """Проходит по строкам листа и наполняет отчёт валидными контактами."""
    seen_inn: dict[str, int] = {}
    first_row = True

    for row_number, row in enumerate(
        worksheet.iter_rows(min_col=COL_NAME + 1, max_col=COL_INN + 1, values_only=True),
        start=1,
    ):
        name_raw = row[COL_NAME] if len(row) > COL_NAME else None
        inn_raw = row[COL_INN] if len(row) > COL_INN else None

        if first_row:
            first_row = False
            if _looks_like_header(name_raw, inn_raw):
                logger.info("Строка %s распознана как заголовок таблицы", row_number)
                continue
            logger.warning(
                "Строка %s не похожа на заголовок (%s / %s) — обрабатываю её как данные",
                row_number,
                _shorten(name_raw, 40),
                _shorten(inn_raw, 40),
            )

        report.total_rows += 1

        name = normalize_name(name_raw)
        if name is None:
            report.skipped_empty_name += 1
            logger.warning("Строка %s: пустое название — пропуск", row_number)
            continue

        inn, reason = normalize_inn(inn_raw)
        if reason == REASON_EMPTY:
            report.skipped_empty_inn += 1
            continue
        if reason is not None:
            report.skipped_invalid_inn += 1
            report.invalid_samples.append(
                f"строка {row_number}: название={_shorten(name, 40)}, "
                f"ИНН={_shorten(inn_raw, 40)}"
            )
            continue
        assert inn is not None  # валидный ИНН всегда строка

        if is_skippable_name(name):
            report.skipped_skip_name += 1
            logger.info(
                "Строка %s: название %r помечено как удаляемое — пропуск",
                row_number,
                name,
            )
            continue

        if inn in seen_inn:
            report.skipped_duplicates += 1
            logger.warning(
                "Строка %s: ИНН %s уже встречался в строке %s — пропуск дубля",
                row_number,
                inn,
                seen_inn[inn],
            )
            continue

        seen_inn[inn] = row_number
        report.contacts.append(
            {
                "company_name": name,
                "inn": inn,
                "email": email,
                "phone": None,
                "is_active": True,
                "notes": f"Импорт из «{source}», строка {row_number}",
            }
        )


def log_report(report: ParseReport, invalid_limit: int = 20) -> None:
    """Печатает сводку по результатам разбора файла."""
    logger.info("Лист: %s", report.sheet_name)
    logger.info("Строк данных: %s", report.total_rows)
    logger.info(
        "Пропущено: пустой ИНН — %s, невалидный ИНН — %s, "
        "название-пометка («удалить») — %s, пустое название — %s, дубликаты ИНН — %s",
        report.skipped_empty_inn,
        report.skipped_invalid_inn,
        report.skipped_skip_name,
        report.skipped_empty_name,
        report.skipped_duplicates,
    )
    logger.info("К импорту: %s контактов (уникальных ИНН)", report.imported)

    if report.invalid_samples and invalid_limit > 0:
        shown = min(invalid_limit, len(report.invalid_samples))
        logger.warning(
            "Невалидных значений ИНН: %s. Первые %s:",
            len(report.invalid_samples),
            shown,
        )
        for sample in report.invalid_samples[:shown]:
            logger.warning("  %s", sample)
        if shown < len(report.invalid_samples):
            logger.warning(
                "  … ещё %s (полный список: --show-invalid %s)",
                len(report.invalid_samples) - shown,
                len(report.invalid_samples),
            )


async def save_contacts(session: AsyncSession, contacts: Sequence[dict[str, Any]]) -> int:
    """Пишет контакты в таблицу contacts, возвращает число обработанных записей.

    Upsert по ИНН (ON CONFLICT (inn) DO UPDATE), поэтому повторный импорт не
    создаёт дублей и обновляет название/email/примечание.
    """
    if not contacts:
        return 0

    dialect = session.bind.dialect.name if session.bind else "postgresql"
    saved = 0

    if dialect == "postgresql":
        for start in range(0, len(contacts), CHUNK_SIZE):
            chunk = list(contacts[start : start + CHUNK_SIZE])
            stmt = pg_insert(Contact).values(chunk)
            stmt = stmt.on_conflict_do_update(
                index_elements=[Contact.inn],
                set_={
                    "company_name": stmt.excluded.company_name,
                    "email": stmt.excluded.email,
                    "phone": stmt.excluded.phone,
                    "is_active": stmt.excluded.is_active,
                    "notes": stmt.excluded.notes,
                },
            )
            await session.execute(stmt)
            saved += len(chunk)
    else:
        # Для SQLite / тестовой БД
        for data in contacts:
            existing = (
                await session.execute(select(Contact).where(Contact.inn == data["inn"]))
            ).scalar_one_or_none()
            if existing:
                existing.company_name = str(data["company_name"])
                existing.email = str(data["email"])
                existing.phone = data.get("phone")  # type: ignore[assignment]
                existing.is_active = bool(data.get("is_active", True))
                existing.notes = data.get("notes")  # type: ignore[assignment]
            else:
                session.add(Contact(**data))  # type: ignore[arg-type]
            saved += 1

    await session.commit()
    logger.info("Записано/обновлено контактов: %s", saved)
    return saved


async def clean_contacts(session: AsyncSession) -> int:
    """Удаляет все записи из таблицы contacts."""
    res = await session.execute(delete(Contact))
    await session.commit()
    count = res.rowcount or 0
    logger.info("Таблица contacts очищена (удалено записей: %s)", count)
    return count


async def main() -> None:
    parser = argparse.ArgumentParser(
        description="Импорт контактов из Excel-файла в таблицу contacts"
    )
    parser.add_argument(
        "--file",
        default=str(DEFAULT_FILE),
        help=f"Путь к файлу Excel (по умолчанию: {DEFAULT_FILE_REL})",
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Имя листа (по умолчанию — первый лист файла)",
    )
    parser.add_argument(
        "--email",
        default=DEFAULT_EMAIL,
        help=f"Email для всех импортируемых контактов (по умолчанию: {DEFAULT_EMAIL})",
    )
    parser.add_argument(
        "--clean",
        action="store_true",
        help="Удалить все записи из таблицы contacts перед импортом",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Только показать отчёт, ничего не писать в БД",
    )
    parser.add_argument(
        "--show-invalid",
        type=int,
        default=20,
        help="Сколько примеров невалидных значений показать (0 — не показывать)",
    )
    args = parser.parse_args()

    try:
        email = validate_email(args.email)
    except ValueError as exc:
        logger.error("%s", exc)
        sys.exit(1)

    path = Path(args.file).expanduser()
    if not path.is_absolute():
        path = (Path.cwd() / path).resolve()

    try:
        report = parse_workbook(path, sheet=args.sheet, email=email)
    except (FileNotFoundError, ValueError) as exc:
        logger.error("%s", exc)
        sys.exit(1)

    logger.info("Файл: %s", path)
    log_report(report, invalid_limit=args.show_invalid)

    if report.imported == 0:
        logger.error("Валидных контактов не найдено — запись в БД не выполняется")
        sys.exit(1)

    if args.dry_run:
        for contact in report.contacts[:3]:
            logger.info(
                "  [dry-run] %s | ИНН %s | %s",
                contact["company_name"],
                contact["inn"],
                contact["email"],
            )
        logger.info("Режим --dry-run: запись в БД пропущена")
        return

    try:
        async with async_session_factory() as session:
            if args.clean:
                await clean_contacts(session)

            count_before = (
                await session.execute(select(func.count()).select_from(Contact))
            ).scalar() or 0
            logger.info("Контактов в базе до импорта: %s", count_before)

            await save_contacts(session, report.contacts)

            count_after = (
                await session.execute(select(func.count()).select_from(Contact))
            ).scalar() or 0
            logger.info("Контактов в базе после импорта: %s", count_after)
    except Exception as exc:  # ошибка БД не должна тонуть в голом трейсбеке
        logger.exception("Ошибка записи в БД: %s", exc)
        sys.exit(1)
    finally:
        await engine.dispose()


if __name__ == "__main__":
    asyncio.run(main())
